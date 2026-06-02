from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.db.database import SessionLocal
from app.schemas.device import (
    DeviceResponse, GenerateTokenRequest, TokenResponse,
    EnrollDeviceRequest, CheckinRequest, CommandCreate, CommandResponse,
    BulkUploadResult, BULK_UPLOAD_COLUMNS,
)
from app.services.device_service import (
    generate_enrollment_token, get_all_tokens, revoke_token,
    enroll_device, get_all_devices, get_device_by_id,
    device_checkin, get_enrollment_history, get_enrollment_stats,
    send_command, get_commands, complete_command, bulk_create_devices,
)
import csv
import io

router = APIRouter()


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# ✅ Enrollment Stats
@router.get("/enrollment/stats")
def enrollment_stats(db: Session = Depends(get_db)):
    return get_enrollment_stats(db)


# ✅ Enrollment History
@router.get("/enrollment/history")
def enrollment_history(db: Session = Depends(get_db)):
    return get_enrollment_history(db)


# ✅ Generate Token
@router.post("/enrollment/tokens", response_model=TokenResponse)
def create_token(data: GenerateTokenRequest, db: Session = Depends(get_db)):
    return generate_enrollment_token(db, data.platform, data.created_by, data.expires_hours)


# ✅ List Tokens
@router.get("/enrollment/tokens", response_model=list[TokenResponse])
def list_tokens(db: Session = Depends(get_db)):
    return get_all_tokens(db)


# ✅ Revoke Token
@router.delete("/enrollment/tokens/{token_id}")
def revoke(token_id: int, db: Session = Depends(get_db)):
    success = revoke_token(db, token_id)
    if not success:
        raise HTTPException(status_code=404, detail="Token not found")
    return {"message": "Token revoked successfully"}


# ✅ Enroll Device
@router.post("/enrollment/enroll")
def enroll(data: EnrollDeviceRequest, db: Session = Depends(get_db)):
    device, error = enroll_device(db, data)
    if error:
        raise HTTPException(status_code=400, detail=error)
    return {
        "device_id": device.id,
        "device_uid": device.device_uid,
        "device_name": device.name,
        "platform": device.platform,
        "enrolled_at": device.last_seen,
        "message": f"Device '{device.name}' enrolled successfully",
    }


# ✅ Device Checkin (heartbeat)
@router.post("/devices/checkin")
def checkin(data: CheckinRequest, db: Session = Depends(get_db)):
    device, error = device_checkin(db, data)
    if error:
        raise HTTPException(status_code=404, detail=error)
    return {"message": "Checkin successful", "last_seen": device.last_seen}


# ✅ List all devices
@router.get("/devices", response_model=list[DeviceResponse])
def list_devices(db: Session = Depends(get_db)):
    return get_all_devices(db)


# ✅ Get device by ID
@router.get("/devices/{device_id}", response_model=DeviceResponse)
def get_device(device_id: int, db: Session = Depends(get_db)):
    device = get_device_by_id(db, device_id)
    if not device:
        raise HTTPException(status_code=404, detail="Device not found")
    return device


# ── Device Commands ───────────────────────────────────────────────────────────

@router.post("/devices/{device_id}/commands", response_model=CommandResponse)
def issue_command(device_id: int, data: CommandCreate, db: Session = Depends(get_db)):
    """Send a command to a device (lock, wipe, reboot, sync_policy, collect_logs)."""
    cmd, error = send_command(db, device_id, data.command)
    if error:
        raise HTTPException(status_code=400, detail=error)
    return cmd


@router.get("/devices/{device_id}/commands", response_model=list[CommandResponse])
def list_commands(device_id: int, db: Session = Depends(get_db)):
    """Get all commands issued to a device."""
    return get_commands(db, device_id)


@router.patch("/devices/{device_id}/commands/{command_id}/complete", response_model=CommandResponse)
def mark_command_complete(device_id: int, command_id: int, db: Session = Depends(get_db)):
    """Mark a pending command as executed (called by device agent on completion)."""
    cmd = complete_command(db, command_id)
    if not cmd:
        raise HTTPException(status_code=404, detail="Command not found")
    return cmd


# ── Bulk Upload ───────────────────────────────────────────────────────────────

@router.get("/devices/bulk-upload/template")
def download_template():
    """Download a CSV template with the correct column headers for bulk device upload."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(BULK_UPLOAD_COLUMNS)
    # Write one example row
    writer.writerow([
        "Laptop-IT-001", "windows", "DEV-ABCD1234", "Dell Latitude 5520", "Dell",
        "SN-123456", "AA:BB:CC:DD:EE:FF", "Windows 11 22H2", "john@risktrix.com",
        "IT Operations", "low", "active",
    ])
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=device_bulk_upload_template.csv"},
    )


@router.post("/devices/bulk-upload", response_model=BulkUploadResult)
async def bulk_upload_devices(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Bulk create devices from a CSV or XLSX file.
    CSV: must have a header row matching the template columns (order does not matter).
    XLSX: first sheet, first row = headers.
    """
    filename = file.filename or ""
    content = await file.read()

    rows: list[dict] = []

    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")  # utf-8-sig strips BOM from Excel-saved CSV
        reader = csv.DictReader(io.StringIO(text))
        rows = [row for row in reader]

    elif filename.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
        except ImportError:
            raise HTTPException(status_code=422, detail="openpyxl not installed. Use CSV format.")
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        headers = [str(cell.value or "").strip().lower() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            rows.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
    else:
        raise HTTPException(status_code=422, detail="Unsupported file type. Upload a .csv or .xlsx file.")

    if not rows:
        raise HTTPException(status_code=400, detail="File is empty or has no data rows.")

    result = bulk_create_devices(db, rows)
    return result